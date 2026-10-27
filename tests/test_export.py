"""Проверка настоящего XLSX и экспорта из текущего состояния Qt-интерфейса.

Запуск: QT_QPA_PLATFORM=offscreen python -m unittest discover -s tests
"""
import csv
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtWidgets import QApplication
from openpyxl import load_workbook

from core.cost import ENERGY_BY_ROUTE
from ui.cost_panel import GROUPS
from ui.export import _save_xlsx, save_table
from ui.pages.route import RoutePage, ROUTE_HEADERS


class TestCostExport(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.page = RoutePage(None)
        self.addCleanup(self.page.deleteLater)
        self.page.safe_recalc()
        self.assertTrue(self.page.rows, self.page.status.text())

    def export(self):
        path = str(Path(self.tmp.name) / "route.xlsx")
        with patch("ui.export.QFileDialog.getSaveFileName", return_value=(path, "")), \
                patch("ui.export.QMessageBox.information"):
            self.page.export()
        wb = load_workbook(path, data_only=True)
        self.addCleanup(wb.close)
        self.assertEqual(wb.sheetnames, ["Волочение", "Себестоимость"])
        return wb

    def values(self, ws):
        return {row[0].value: row[1] for row in ws if row[0].value is not None}

    def test_current_inputs_results_and_numeric_cells(self):
        p = self.page
        p.cost_full.setChecked(True)
        p.cost_copper.setChecked(True)
        p.cost_method.setCurrentText(ENERGY_BY_ROUTE)
        p.cost_f["tariff"].setText("12,3456")
        p.cost_f["cu_um"].setText("0,45")
        r = p.res["cost"]
        self.assertIsNotNone(r)
        wb = self.export()
        ws = wb["Себестоимость"]
        values = self.values(ws)
        for _, fields in GROUPS:
            for key, title, _ in fields:
                self.assertEqual(values[title].data_type, "n", title)
                self.assertAlmostEqual(values[title].value,
                                       float(p.cost_f[key].text().replace(",", ".")))
        self.assertEqual(values["Метод расчёта электроэнергии"].value, ENERGY_BY_ROUTE)
        self.assertAlmostEqual(values["ИТОГО передел"].value, r["conversion"])
        self.assertAlmostEqual(values["Передел, руб/кг"].value, r["conversion"] / 1000)
        self.assertAlmostEqual(values["ИТОГО полная себестоимость"].value, r["full"])
        for key, title in [("q_t", "Теоретическая производительность, кг/ч"),
                           ("q_f", "Фактическая производительность, кг/ч"),
                           ("kwh_t", "Электроэнергия, кВт·ч/т"),
                           ("p_work", "Рабочая мощность, кВт"),
                           ("out_year_t", "Годовой выпуск, т"), ("cu_kg_t", "Медь, кг/т")]:
            self.assertAlmostEqual(values[title].value, r[key])
        exported = {row[0]: row[1:] for row in ws.iter_rows(values_only=True)}
        for name, value in r["items"]:
            self.assertAlmostEqual(exported[name][0], value)
            self.assertAlmostEqual(exported[name][1], value / r["conversion"] * 100)
        for name, value in r["full_items"]:
            self.assertAlmostEqual(exported[name][0], value)
            self.assertIsNone(exported[name][1])
        # Existing route headers, numeric data and highlighting survive.
        route = wb["Волочение"]
        header = next(row[0].row for row in route if row[0].value == "№")
        self.assertEqual([c.value for c in route[header]], ROUTE_HEADERS)
        self.assertEqual(route.cell(header + 1, 2).value, round(p.rows[0]["d_in"], 4))
        for i, row in enumerate(p.rows, header + 1):
            if row["anneal"]:
                self.assertEqual(route.cell(i, 1).fill.fgColor.rgb, "00FFF0C2")

    def test_disabled_preserves_inputs_without_stale_totals(self):
        self.page.cost_on.setChecked(False)
        self.page.cost_f["tariff"].setText("9,75")
        values = self.values(self.export()["Себестоимость"])
        self.assertFalse(values["Расчёт себестоимости включён"].value)
        self.assertEqual(values["Тариф, руб/кВт·ч"].value, 9.75)
        self.assertIn("выключен", values["Статус"].value)
        self.assertNotIn("ИТОГО передел", values)

    def test_invalid_cost_does_not_prevent_route_export(self):
        self.page.cost_f["tariff"].setText("ошибка")
        values = self.values(self.export()["Себестоимость"])
        self.assertEqual(values["Тариф, руб/кВт·ч"].value, "ошибка")
        self.assertIn("не рассчитана", values["Статус"].value)
        self.assertNotIn("ИТОГО передел", values)

    def test_optional_costs_off_and_live_edit(self):
        before = self.page.res["cost"]["conversion"]
        self.page.cost_f["tariff"].setText("20,1")
        values = self.values(self.export()["Себестоимость"])
        self.assertNotEqual(values["ИТОГО передел"].value, before)
        self.assertAlmostEqual(values["ИТОГО передел"].value,
                               self.page.res["cost"]["conversion"])
        self.assertNotIn("ИТОГО полная себестоимость", values)
        self.assertNotIn("Медь на омеднение", values)
        # Inactive optional inputs are still exported.
        self.assertEqual(values["Цена катанки, руб/т"].value, 85000)

    def test_shared_export_defaults_and_csv_remain_compatible(self):
        path = str(Path(self.tmp.name) / "generic.xlsx")
        _save_xlsx(path, [("Параметр", 1.5)], ["Значение"], [[2.5]], ["red"])
        wb = load_workbook(path)
        self.addCleanup(wb.close)
        self.assertEqual(wb.sheetnames, ["Таблица"])
        self.assertEqual(wb.active["A4"].value, 2.5)
        self.assertEqual(wb.active["A4"].fill.fgColor.rgb, "00FBE4E2")
        path = str(Path(self.tmp.name) / "generic.csv")
        with patch("ui.export.QFileDialog.getSaveFileName", return_value=(path, "")), \
                patch("ui.export.QMessageBox.information"):
            save_table(None, "test", [], ["Значение"], [[2.5]],
                       extra_sheets=[("Себестоимость", [["ignored"]])])
        with open(path, encoding="utf-8-sig", newline="") as f:
            self.assertEqual(list(csv.reader(f, delimiter=";")), [["Значение"], ["2,5"]])


if __name__ == "__main__":
    unittest.main()
