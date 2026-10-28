"""Сохранение таблиц в Excel (.xlsx) или CSV — общий код для всех разделов."""
import csv
import os

from PySide6.QtWidgets import QFileDialog, QMessageBox

from ui.popup import notify

FILLS = {"yellow": "FFF0C2", "red": "FBE4E2"}


def save_table(parent, default_name, info, headers, rows, fills=None):
    """info — пары (параметр, значение) над таблицей; rows — строки таблицы;
    fills — подсветка строк: 'yellow', 'red' или None."""
    start = os.path.join(os.path.expanduser("~"), default_name + ".xlsx")
    path, _ = QFileDialog.getSaveFileName(parent, "Сохранить таблицу", start,
                                          "Excel (*.xlsx);;CSV для Excel (*.csv)")
    if not path:
        return
    fills = fills or [None] * len(rows)
    try:
        if path.lower().endswith(".xlsx"):
            try:
                _save_xlsx(path, info, headers, rows, fills)
            except ImportError:
                QMessageBox.warning(parent, "Нужен openpyxl",
                                    "Для файлов Excel установи пакет командой\n"
                                    "pip install openpyxl\n\nСейчас сохраню в CSV.")
                path = os.path.splitext(path)[0] + ".csv"
                _save_csv(path, info, headers, rows)
        else:
            _save_csv(path, info, headers, rows)
        notify(parent, "Файл сохранён", "ok", detail=path)
    except OSError as e:
        QMessageBox.warning(parent, "Не удалось сохранить", str(e))


def save_workbook(parent, default_name, sheets):
    """Книга Excel из нескольких листов: sheets — [(название, info, headers, rows, fills)].
    Возвращает путь или None. Без openpyxl — один CSV, листы друг под другом."""
    start = os.path.join(os.path.expanduser("~"), default_name + ".xlsx")
    path, _ = QFileDialog.getSaveFileName(parent, "Сохранить в Excel", start,
                                          "Excel (*.xlsx);;CSV для Excel (*.csv)")
    if not path:
        return None
    try:
        if path.lower().endswith(".xlsx"):
            try:
                import openpyxl
            except ImportError:
                path = os.path.splitext(path)[0] + ".csv"
            else:
                wb = openpyxl.Workbook()
                wb.remove(wb.active)
                for title, info, headers, rows, fills in sheets:
                    _fill_sheet(wb.create_sheet(_sheet_title(title)), info, headers, rows,
                                fills or [None] * len(rows))
                wb.save(path)
                return path
        with open(path, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh, delimiter=";")
            for title, info, headers, rows, _ in sheets:
                w.writerow([title])
                _write_csv(w, info, headers, rows)
                w.writerow([])
        return path
    except OSError as e:
        QMessageBox.warning(parent, "Не удалось сохранить", str(e))
        return None


def _sheet_title(title):
    """Название листа: до 31 символа, без символов, которые Excel не принимает."""
    for ch in "[]:*?/\\":
        title = title.replace(ch, " ")
    return title[:31]


def _save_xlsx(path, info, headers, rows, fills):
    import openpyxl
    wb = openpyxl.Workbook()
    _fill_sheet(wb.active, info, headers, rows, fills)
    wb.active.title = "Таблица"
    wb.save(path)


def _fill_sheet(ws, info, headers, rows, fills):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    for k, val in info:
        ws.append([k, val])
    if info:
        ws.append([])
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")   # длинный заголовок — в 2 строки
    for row, fill in zip(rows, fills):
        ws.append(row)
        if fill in FILLS:
            pf = PatternFill("solid", fgColor=FILLS[fill])
            for cell in ws[ws.max_row]:
                cell.fill = pf
    # ширина — по самому длинному слову заголовка и значениям столбца
    for i, h in enumerate(headers, 1):
        words = max((len(w) for w in str(h).split()), default=8)
        vals = max((len(str(row[i - 1])) for row in rows if i - 1 < len(row)), default=0)
        width = min(max(10, words + 2, vals + 2), 60)
        if i == 1 and info:
            width = max(width, min(max(len(str(k)) for k, _ in info) + 2, 60))
        ws.column_dimensions[get_column_letter(i)].width = width


def _save_csv(path, info, headers, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        _write_csv(csv.writer(fh, delimiter=";"), info, headers, rows)


def _write_csv(w, info, headers, rows):
    def cell(x):
        return str(x).replace(".", ",") if isinstance(x, float) else x
    for k, val in info:
        w.writerow([k, cell(val)])
    if info:
        w.writerow([])
    w.writerow(headers)
    for row in rows:
        w.writerow([cell(x) for x in row])


def table_sheet(title, table):
    """Лист из таблицы на экране: заголовки и ячейки как есть, числа — числами."""
    heads = [table.horizontalHeaderItem(c).text().replace("\n", " ")
             for c in range(table.columnCount())]
    rows = []
    for r in range(table.rowCount()):
        row = []
        for c in range(table.columnCount()):
            it = table.item(r, c)
            txt = it.text() if it else ""
            try:
                row.append(float(txt.replace(",", ".")) if txt not in ("", "—") else txt)
            except ValueError:
                row.append(txt)
        rows.append(row)
    return (title, [], heads, rows, None)
